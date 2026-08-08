from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from icalendar import Calendar, Event
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt

from app.config import get_settings
from app.db.session import db_conn
from app.security.crypto import hmac_sign, sha256_hex


def _outdir(run_id: str) -> Path:
    settings = get_settings()
    root = settings.resolve_path(settings.artifact_temp_dir) / run_id
    root.mkdir(parents=True, exist_ok=True)
    return root


def _register(
    run_id: str,
    *,
    artifact_type: str,
    path: Path,
    content_type: str,
    manifest: dict[str, Any] | None = None,
) -> dict[str, Any]:
    data = path.read_bytes()
    digest = sha256_hex(data)
    signature = hmac_sign(digest)
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO tuntas.artifacts
                  (run_id, artifact_type, filename, storage_path, content_type,
                   sha256, hmac_signature, byte_size, manifest)
                VALUES (%s::uuid, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                RETURNING id::text, filename, sha256, hmac_signature, byte_size, artifact_type
                """,
                (
                    run_id,
                    artifact_type,
                    path.name,
                    str(path),
                    content_type,
                    digest,
                    signature,
                    len(data),
                    json.dumps(manifest or {}),
                ),
            )
            row = dict(cur.fetchone())
        conn.commit()
    return row


def export_management_pack(
    run_id: str,
    *,
    trigger: dict[str, Any],
    portfolios: list[dict[str, Any]],
    selected: dict[str, Any] | None,
    sessions: list[dict[str, Any]],
    secretariat: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    out = _outdir(run_id)
    artifacts: list[dict[str, Any]] = []

    # 1) DOCX approval paper
    doc = Document()
    doc.add_heading("TUNTAS Management Approval Pack", 0)
    req = trigger["request"]
    doc.add_paragraph(f"Request ID: {req.get('request_id')}")
    doc.add_paragraph(f"Issuing Unit: {req.get('issuing_unit')}")
    doc.add_paragraph(f"Capability Domain: {req.get('capability_domain')}")
    doc.add_paragraph(
        (secretariat or {}).get("decision_brief")
        or "AI-prepared brief pending manager decision."
    )
    doc.add_heading("Portfolio Options", level=1)
    for p in portfolios:
        doc.add_paragraph(
            f"{p['label']}: MYR {p['total_cost_myr']} | coverage={p['coverage_score']} | "
            f"hard_ok={p['hard_constraint_ok']}",
            style="List Bullet",
        )
    if selected:
        doc.add_heading("Selected Option", level=1)
        doc.add_paragraph(selected.get("label", selected.get("option_key", "")))
    docx_path = out / "management_approval.docx"
    doc.save(docx_path)
    artifacts.append(
        _register(
            run_id,
            artifact_type="docx",
            path=docx_path,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    )

    # 2) XLSX vendor / assignment comparison
    wb = Workbook()
    ws = wb.active
    ws.title = "Assignments"
    ws.append(
        ["option_key", "employee_ref", "course_code", "course_title", "provider_code", "cost_myr"]
    )
    for p in portfolios:
        for a in p.get("assignments") or []:
            ws.append(
                [
                    p["option_key"],
                    a.get("employee_ref"),
                    a.get("course_code"),
                    a.get("course_title"),
                    a.get("provider_code"),
                    a.get("cost_myr"),
                ]
            )
    ws2 = wb.create_sheet("Options")
    ws2.append(
        [
            "option_key",
            "label",
            "total_cost_myr",
            "cost_per_employee_myr",
            "coverage_score",
            "hard_constraint_ok",
        ]
    )
    for p in portfolios:
        ws2.append(
            [
                p["option_key"],
                p["label"],
                p["total_cost_myr"],
                p["cost_per_employee_myr"],
                p["coverage_score"],
                p["hard_constraint_ok"],
            ]
        )
    xlsx_path = out / "vendor_portfolio_comparison.xlsx"
    wb.save(xlsx_path)
    artifacts.append(
        _register(
            run_id,
            artifact_type="xlsx",
            path=xlsx_path,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    )

    # 3) PPTX exec summary
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(1))
    tf = title_box.text_frame
    tf.text = "TUNTAS Capability Assurance Brief"
    tf.paragraphs[0].font.size = Pt(28)
    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    bf = body.text_frame
    bf.text = f"{req.get('request_id')} — {req.get('capability_domain')}"
    for p in portfolios:
        para = bf.add_paragraph()
        para.text = f"{p['label']}: MYR {p['total_cost_myr']} | cov {p['coverage_score']}"
        para.level = 1
    pptx_path = out / "exec_brief.pptx"
    prs.save(pptx_path)
    artifacts.append(
        _register(
            run_id,
            artifact_type="pptx",
            path=pptx_path,
            content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
    )

    # 4) JSON evidence manifest
    manifest = {
        "run_id": run_id,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "request_id": req.get("request_id"),
        "portfolios": [
            {
                "option_key": p["option_key"],
                "total_cost_myr": p["total_cost_myr"],
                "hard_constraint_ok": p["hard_constraint_ok"],
            }
            for p in portfolios
        ],
        "selected_option_key": (selected or {}).get("option_key"),
        "session_count": len(sessions),
    }
    json_path = out / "evidence_manifest.json"
    json_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    artifacts.append(
        _register(
            run_id,
            artifact_type="json",
            path=json_path,
            content_type="application/json",
            manifest=manifest,
        )
    )

    # 5) ICS calendar
    cal = Calendar()
    cal.add("prodid", "-//TUNTAS//Capability Schedule//EN")
    cal.add("version", "2.0")
    for s in sessions:
        ev = Event()
        ev.add("summary", s["title"])
        ev.add("dtstart", datetime.fromisoformat(s["starts_at"]))
        ev.add("dtend", datetime.fromisoformat(s["ends_at"]))
        ev.add("location", s.get("location") or "")
        ev.add("description", f"Course {s['course_code']} via TUNTAS")
        cal.add_component(ev)
    ics_path = out / "training_schedule.ics"
    ics_path.write_bytes(cal.to_ical())
    artifacts.append(
        _register(
            run_id,
            artifact_type="ics",
            path=ics_path,
            content_type="text/calendar",
        )
    )

    # 6) PDF via WeasyPrint — required (no text fallback)
    from weasyprint import HTML

    selected_label = (selected or {}).get("label") or (selected or {}).get("option_key") or "n/a"
    html = f"""
    <html>
    <head>
      <meta charset="utf-8"/>
      <style>
        body {{ font-family: Helvetica, Arial, sans-serif; margin: 48px; color: #111; }}
        h1 {{ font-size: 22px; margin-bottom: 8px; }}
        h2 {{ font-size: 16px; margin-top: 28px; }}
        .meta {{ color: #444; font-size: 12px; }}
        li {{ margin: 6px 0; }}
        .badge {{ display: inline-block; padding: 2px 8px; background: #eef6ff; border: 1px solid #bcd; }}
      </style>
    </head>
    <body>
      <h1>TUNTAS Assurance Pack</h1>
      <p class="meta">Traceable Upskilling &amp; Normative Training Assurance System</p>
      <p><span class="badge">Request {req.get('request_id')}</span></p>
      <p>Capability domain: {req.get('capability_domain')}</p>
      <p>Planned proficiency movement: Level {req.get('proficiency_from')} → Level {req.get('proficiency_to')}</p>
      <p>Selected portfolio: <strong>{selected_label}</strong></p>
      <h2>Portfolio options (OR-Tools CP-SAT)</h2>
      <ul>
      {''.join(
          f"<li>{p['label']}: MYR {p['total_cost_myr']} | coverage={p['coverage_score']} | hard_ok={p['hard_constraint_ok']}</li>"
          for p in portfolios
      )}
      </ul>
      <h2>Schedule sessions</h2>
      <ul>
      {''.join(
          f"<li>{s.get('title')} — {s.get('starts_at')} ({len(s.get('employee_refs') or [])} staff)</li>"
          for s in sessions
      ) or '<li>Pending approval / no sessions committed</li>'}
      </ul>
      <p class="meta">Generated: {datetime.now(timezone.utc).isoformat()} · Evidence-backed decision support — not legal advice.</p>
    </body>
    </html>
    """
    pdf_path = out / "assurance_pack.pdf"
    HTML(string=html).write_pdf(str(pdf_path))
    if not pdf_path.exists() or pdf_path.stat().st_size < 100:
        raise RuntimeError("WeasyPrint failed to produce a valid PDF")
    artifacts.append(
        _register(
            run_id,
            artifact_type="pdf",
            path=pdf_path,
            content_type="application/pdf",
        )
    )

    return artifacts


def list_artifacts(run_id: str) -> list[dict[str, Any]]:
    with db_conn() as conn:
        with conn.cursor() as cur:
            cur.execute(
                """
                SELECT id::text, artifact_type, filename, storage_path, content_type,
                       sha256, hmac_signature, byte_size, created_at
                FROM tuntas.artifacts
                WHERE run_id = %s::uuid
                ORDER BY created_at
                """,
                (run_id,),
            )
            return [dict(r) for r in cur.fetchall()]
